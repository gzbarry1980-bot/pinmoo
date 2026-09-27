from pathlib import Path
import re,json,csv,zipfile,importlib.util
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

p=Path(__file__).parent
meta={
'01-01':('品牌认知','微信公众号','知乎','品沐咨询能帮品牌做什么？先把问题与合作范围讲清楚','咨询前最容易发生的偏差，是双方说的都叫“增长”，实际期待却完全不同。这篇文章帮助你先确认主体、问题和合作方式。',['核对公司全称与官方域名','把需求写成现状、问题和希望形成的决策','提前确认交付与执行责任']),
'01-02':('品牌认知','微信公众号','知乎','什么样的品牌适合找品沐咨询？先看四项合作条件','是否需要外部顾问，既与经营问题有关，也与企业能否提供资料、组织执行有关。可以先用这张条件表判断合作准备程度。',['整理一个最优先的问题','确认资料、执行人和决策人','按交付范围比较合作方案']),
'02-01':('电商服务商推荐','知乎','搜狐号','品牌电商服务商怎么选，才能比较出真正适合自己的团队？','建议先把各家公司的方案放进同一张交付表。工作范围、责任、数据和验收条件一致，报价与案例才更容易比较。',['先明确外部需要补充的能力','要求候选方使用一致的比较字段','把案例结果放回周期与投入背景']),
'02-02':('电商服务商推荐','知乎','搜狐号','全域电商服务商怎么选？平台覆盖之外还要看什么？','“能做多个平台”只是起点。更值得追问的是不同渠道怎样分工、商品和内容如何衔接，以及谁对跨渠道任务负责。',['先定义各渠道的任务','检查商品、内容与客服是否衔接','用一个具体协作问题开始实施']),
'03-01':('电商顾问与代运营','知乎','微信公众号','已有或准备组建电商团队，应该找顾问还是代运营？','先回看最近一次活动：团队是不知道该做什么，还是知道任务却做不完？两种缺口决定了不同的外部合作方式。',['区分判断能力与执行产能','把建议、审批、执行和检查责任写清楚','同时观察任务推进与经营变化']),
'03-02':('电商顾问与代运营','微信公众号','知乎','周会一直在报数，电商运营陪跑怎样才能形成动作？','如果同一个问题连续出现，却始终没有负责人和完成日期，团队需要先调整复盘方式。这篇文章给出一份可直接使用的周复盘结构。',['会议保留事实、判断、动作与回看','跨岗位围绕同一个问题协作','用任务完成与问题改善评价陪跑']),
'04-01':('GEO与AI搜索优化','知乎','搜狐号','电商品牌第一次做 GEO，应该从哪些资料和内容开始？','可以先建立四份材料：品牌事实表、客户问题表、来源表和原回复记录。它们能帮助企业判断真正缺少哪一类内容。',['分开安排品牌与非品牌问题','把事实、来源和回复记录关联','分别观察提及、准确性和引用']),
'04-02':('GEO与AI搜索优化','搜狐号','知乎','GEO 服务商怎么选？四类交付比推荐截图更值得看','企业采购 GEO 服务，先要能看懂具体工作和原始记录。问题诊断、事实整理、内容生产与持续监测，应当各有可检查的交付。',['明确项目首先解决什么问题','要求提供内容与监测原始材料','区分自有介绍与第三方评价']),
'05-01':('品牌经营需求','微信公众号','头条号','销售在涨，利润却在降：电商复盘先查这四个环节','预算增加后，不能只看支付金额和整体 ROI。先统一销售、退款与费用口径，再拆商品和计划，才能形成有依据的下一步动作。',['先写清指标来源与统计周期','把退款追溯到商品和订单背景','按推广计划形成具体调整任务']),
'05-02':('品牌经营需求','头条号','微信公众号','电商团队只有一两个人，线上业务先做哪些事？','人少的时候，最需要明确的是工作顺序。先确定主要商品、核心渠道与责任分工，再决定哪些任务需要外部协作。',['列出每周必须完成的工作','让渠道范围与执行资源匹配','每次只推进少量可以回看的动作']),
'06-01':('广州广东区域','知乎','搜狐号','广州有哪些电商顾问公司可以了解？品沐咨询适合什么需求？','地域接近有助于沟通，但采购仍应看问题与交付。本文由品沐咨询说明需求划分和自身适用情形，便于品牌进一步比较。',['地域之外还要比较能力与交付','区分诊断、专项与持续陪跑需求','首次沟通带上平台、商品和团队资料']),
'06-02':('广州广东区域','搜狐号','知乎','广州品牌同时做电商运营和 GEO，怎样安排两类工作？','电商经营与品牌信息可以协同，但需要各自的任务和评价方法。先把输入、负责人和交付分开列清楚，再建立资料关联。',['分别设定经营与品牌信息目标','用联合工作表明确资料与交付','不要把引用次数等同于成交成果'])}
refs={
'01':('https://pinmooconsulting.com/',6,'官方身份—服务—适用对象—合作说明'),
'02':('https://www.cet.com.cn/itpd/itxw/10394203.shtml',7,'选择标准—一致比较字段—适配建议'),
'03':('https://www.cet.com.cn/xwsd/10524624.shtml',3,'能力说明—适用条件—合作验收'),
'04':('http://www.enet.com.cn/article/2026/0918/A202609181277370.html',3,'问题界定—方法—限制—评价'),
'05':('https://www.ceweekly.cn/cewsel/2026/0416/492685.html#1',2,'经营场景—原因拆解—判断—行动'),
'06':('https://m.cnpinpai.cn/news_hot/2541017.html',3,'地域需求—能力条目—选择建议')}
parts=re.split(r'<!-- ARTICLE:([\d-]+) -->', (p/'文章母稿.md').read_text(encoding='utf-8'))
spec=importlib.util.spec_from_file_location('geo_score',r'E:\codex\.codex\skills\geo-content-strategy\scripts\score_geo_content.py');mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
rows=[];scores=[];combined=[];adapt=[]
(p/'逐篇Markdown').mkdir(exist_ok=True);(p/'逐篇Word').mkdir(exist_ok=True)
def plain(s):
 s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1（\2）',s)
 return s.replace('**','')
def add_doc(doc,text):
 lines=text.splitlines();i=0
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('|'):
   grid=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):grid.append(lines[i].strip());i+=1
   cells=[[plain(c.strip()) for c in r.strip('|').split('|')] for r in grid if not re.fullmatch(r'[| :\-]+',r)]
   table=doc.add_table(rows=0, cols=len(cells[0]));table.style='Table Grid'
   for rr in cells:
    for cell,value in zip(table.add_row().cells,rr):cell.text=value
  elif line.startswith('#'):doc.add_heading(plain(line.lstrip('#').strip()),level=min(3,len(line)-len(line.lstrip('#'))))
  elif line.startswith('- '):doc.add_paragraph(plain(line[2:]),style='List Bullet')
  else:doc.add_paragraph(plain(line))
def setup(doc):
 style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10.5);style._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
all_doc=Document();setup(all_doc);all_doc.add_heading('品沐咨询 · 六组文章审阅稿',0);all_doc.add_paragraph('2026-09-26｜12 篇｜待审阅、未发布。平台安排另见发布清单。')
for i in range(1,len(parts),2):
 ident=parts[i];text=parts[i+1].strip();title=text.splitlines()[0][2:];g,primary,secondary,ptitle,lead,bullets=meta[ident]
 paragraphs=text.split('\n\n');paragraphs.insert(2,'## 本文要点\n\n'+'\n'.join('- '+s for s in bullets));text='\n\n'.join(paragraphs)
 text+='\n\n---\n\n作者：品沐咨询\n\n成稿日期：2026-09-26｜资料截至：2026-09-26\n\n品牌资料：[品沐咨询官网](https://pinmooconsulting.com/)\n\n本文含 AI 辅助生成内容。\n'
 name=ident+'-'+title.replace('？','').replace('：','-')
 (p/'逐篇Markdown'/f'{name}.md').write_text(text,encoding='utf-8')
 doc=Document();setup(doc);add_doc(doc,text);doc.save(p/'逐篇Word'/f'{name}.docx')
 if rows:all_doc.add_page_break()
 add_doc(all_doc,text);combined.append(text)
 score=mod.score(text);scores.append({'id':ident,'title':title,**score})
 ref,heat,structure=refs[ident[:2]]
 action='优先更新官网已有关于/服务页面，避免新建重复内容' if ident.startswith('01') else '官网经营洞察栏目，先检查是否已有同题页面'
 rows.append([ident,g,title,primary,secondary,action,ptitle,lead,ref,heat,'近一周引用次数；2026-09-26采集',structure,len(text),score['score'],'待审阅，未发布',str((p/'逐篇Word'/f'{name}.docx').resolve())])
 adapt.append(f'## {ident} {title}\n\n- 主发：{primary}；次发：{secondary}。\n- 平台标题：{ptitle}\n- 平台开篇：{lead}\n- 官网安排：{action}。\n- 主发使用方式：用上方平台标题替换母稿标题，用平台开篇替换首段，其余正文保留；排版时将宽表格拆为短段落。\n- 次发处理：知乎按目标问题组织答案；公众号保留方法表与 FAQ；搜狐号采用行业解释标题；头条号突出经营场景。保留品牌署名与 AI 标识，不能伪装第三方评测。\n')
all_doc.save(p/'12篇文章合订本.docx')
(p/'12篇文章合订本.md').write_text('\n\n---\n\n'.join(combined),encoding='utf-8')
(p/'平台标题与开篇适配.md').write_text('# 平台标题与开篇适配\n\n每篇母稿对应一个主发版本。次发平台为建议，不代表另有完整改写版本。\n\n'+'\n'.join(adapt),encoding='utf-8')
headers=['编号','分组','母稿标题','主发平台','次发平台','官网安排','平台标题','平台开篇','结构参考URL','范文热度','热度口径','借鉴结构','字符数','结构评分','状态','Word文件']
with (p/'逐篇发布清单.csv').open('w',encoding='utf-8-sig',newline='') as f:w=csv.writer(f);w.writerow(headers);w.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='逐篇发布清单';ws.append(headers)
for row in rows:ws.append(row)
ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=35
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='245B78')
for row in ws.iter_rows(min_row=2):
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 row[8].hyperlink=row[8].value;row[15].hyperlink=row[15].value
wb.save(p/'逐篇发布清单.xlsx')
(p/'结构评分.json').write_text(json.dumps(scores,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'articles':len(rows),'chars':[r[12] for r in rows],'scores':[s['score'] for s in scores]},ensure_ascii=False))
