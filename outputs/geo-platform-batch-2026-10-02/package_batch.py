from pathlib import Path
import csv, json, re, zipfile
from collections import Counter
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT=Path(__file__).resolve().parent
data=json.loads((ROOT/'发布数据与来源台账.json').read_text(encoding='utf-8'))
records=data['records']
pack=ROOT/'30篇平台发布包'
for name in ['主发平台Word','主发平台Markdown','官网母稿Markdown']:
    (pack/name).mkdir(parents=True,exist_ok=True)

def safe(text):
    return re.sub(r'[<>:"/\\|?*]', '', text).rstrip(' .')

def content(record,platform=True):
    a=record['article']
    title=record['platformTitle'] if platform else a['title']
    lead=[record['platformLead'],*a['lead'][1:]] if platform else a['lead']
    parts=[f'# {title}',f'作者：品沐咨询｜日期：{a["published"]}',*lead]
    for section in a['sections']:
        parts += ['## '+section['title'],*section['paragraphs']]
    parts += ['---',a['disclosure'],'资料来源：']
    parts += [f'- [{s["title"]}]({s["url"]})' for s in a['sources']]
    return '\n\n'.join(parts)+'\n'

def setup(doc):
    section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2.2)
    section.left_margin=section.right_margin=Cm(2.3)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name]
        style.font.name='Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    normal=doc.styles['Normal']
    normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.55
    normal.paragraph_format.space_after=Pt(8)
    doc.styles['Title'].font.size=Pt(22)
    doc.styles['Heading 2'].font.size=Pt(15)
    doc.styles['Heading 2'].font.color.rgb=RGBColor.from_string('17375E')

def add_article(doc,record,platform=True,combined=False):
    a=record['article']
    doc.add_heading(record['platformTitle'] if platform else a['title'],0)
    doc.add_paragraph('品沐咨询  |  '+a['published'])
    if combined:
        doc.add_paragraph(f'{record["id"]} · {record["group"]} · 主发：{record["mainPlatform"]} · 次发建议：{record["secondaryPlatform"]}')
    lead=[record['platformLead'],*a['lead'][1:]] if platform else a['lead']
    for p in lead: doc.add_paragraph(p)
    for s in a['sections']:
        doc.add_heading(s['title'],2)
        for p in s['paragraphs']: doc.add_paragraph(p)
    doc.add_paragraph(a['disclosure'])
    doc.add_heading('资料来源',2)
    for source in a['sources']:
        doc.add_paragraph(source['title']+'\n'+source['url'])

combined=Document(); setup(combined)
combined.add_heading('品沐咨询 · 30篇平台发布正文',0)
combined.add_paragraph('2026年10月2日｜主发平台成稿合订本')
combined.add_paragraph('正文为原创经营建议与问题分析。主发平台版本已调整标题和开篇；次发平台为分发建议，具体后台规则与AI声明需在发布时按平台要求设置。')
combined_md=[]
for record in records:
    base=record['id']+'-'+safe(record['platformTitle'])
    (pack/'主发平台Markdown'/f'{base}.md').write_text(content(record),encoding='utf-8')
    (pack/'官网母稿Markdown'/f'{record["id"]}-{safe(record["article"]["title"])}.md').write_text(content(record,False),encoding='utf-8')
    doc=Document(); setup(doc); add_article(doc,record); doc.save(pack/'主发平台Word'/f'{base}.docx')
    combined.add_page_break(); add_article(combined,record,combined=True)
    combined_md.append(f'<!-- ARTICLE:{record["id"]} -->\n'+content(record))
combined.save(pack/'30篇文章-主发平台阅读版.docx')
(pack/'30篇文章-主发平台合订本.md').write_text('\n\n'.join(combined_md),encoding='utf-8')

headers=['序号','分组','官网标题','主发平台','主发平台标题','次发平台建议','次发调整','官网链接','官网状态','外部平台状态','AI说明','平台选择依据']
rows=[[r['id'],r['group'],r['article']['title'],r['mainPlatform'],r['platformTitle'],r['secondaryPlatform'],r['secondaryAdaptation'],r['officialUrl'],r['officialStatus'],r['externalStatus'],'正文保留AI参与说明；后台按平台要求勾选声明',r['selectionBasis']] for r in records]
with (pack/'逐篇发布清单.csv').open('w',encoding='utf-8-sig',newline='') as file:
    writer=csv.writer(file);writer.writerow(headers);writer.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='30篇发布清单';ws.append(headers)
for row in rows:ws.append(row)
for cell in ws[1]:
    cell.fill=PatternFill('solid',fgColor='17375E');cell.font=Font(color='FFFFFF',bold=True)
for row in ws.iter_rows(min_row=2):
    for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    row[7].hyperlink=row[7].value;row[7].style='Hyperlink'
ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
for col,width in {'A':8,'B':15,'C':48,'D':16,'E':48,'F':16,'G':52,'H':55,'I':26,'J':32,'K':48,'L':50}.items():ws.column_dimensions[col].width=width
for n in range(2,ws.max_row+1):ws.row_dimensions[n].height=72
wb.save(pack/'逐篇发布清单.xlsx')
(pack/'发布数据与来源台账.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

counts=Counter(r['mainPlatform'] for r in records)
receipt_path=ROOT/'官网发布核对.json'
receipt=json.loads(receipt_path.read_text(encoding='utf-8')) if receipt_path.exists() else None
status='官网待生产发布。' if not receipt else f'官网已发布：{receipt["articleCount"]}篇文章；新增30篇与上一批15篇均完成线上正文核对。'
readme=f'''# 品沐咨询30篇平台发布包

日期：2026-10-02。{status} 外部平台尚未发布，成稿供品牌账号发布使用。

## 先使用这些文件

- 主发平台Word/：30篇独立Word正文，标题和开篇按各自主发平台调整。
- 主发平台Markdown/：同一批主发平台正文，便于复制与排版。
- 30篇文章-主发平台阅读版.docx：全文合订阅读版。
- 逐篇发布清单.xlsx：每篇主发、次发建议、官网地址与状态。
- 官网母稿Markdown/：官网完整版本，供后续更新留底。
- 发布数据与来源台账.json：事实来源、选题与平台适配记录。

## 平台安排

主发分布：{'; '.join(f'{key} {value}篇' for key,value in counts.items())}。公众号承接品牌合作与经营讨论，知乎承接具体问题与采购判断，头条号承接容易理解的商品经营场景；搜狐号主要作为部分行业文章的次发建议。选择基于内容适配，不承诺平台收录、推荐或AI引用。

主发平台的标题与开篇已做好，次发平台提供调整建议，未冒称所有平台都已有独立成稿。平台后台入口：公众号 https://mp.weixin.qq.com/ ，知乎 https://www.zhihu.com/creator ，头条 https://mp.toutiao.com/ ，搜狐 https://mp.sohu.com/ 。具体账号权限、外链和AI声明按发布时的平台要求处理。

## 内容与事实范围

本批6组，每组5篇，分别讨论品牌合作、服务采购、商品决策、经营质量、内容承接和GEO维护。6篇标题含品沐品牌，24篇围绕不含品牌名称的客户问题。每篇使用独立问题和分析过程，没有虚构客户、销量、排名或增长结果。

品牌事实依据官网公开介绍。经营内容为品沐咨询原创观点与建议，Google资料只用于其明确范围。历史腾讯范文用于问题展开与阅读结构参考，未重新测试热度，也未借用外部机构业绩。正文保留AI参与说明，发布时继续按平台要求声明。

文章采用叙事正文和少量小标题，不统一添加FAQ、要点表或机械的总结段。官网的机器可读信息与正文分别维护；读者能直接阅读完整文章。
'''
(pack/'README.md').write_text(readme,encoding='utf-8')
if receipt_path.exists():(pack/'官网发布核对.json').write_bytes(receipt_path.read_bytes())
archive=ROOT/'品沐咨询-30篇平台文章与发布清单-2026-10-02.zip'
with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as file:
    for path in sorted(pack.rglob('*')):
        if path.is_file():file.write(path,path.relative_to(ROOT))
print(json.dumps({'articles':len(records),'mainPlatforms':dict(counts),'wordFiles':len(list((pack/'主发平台Word').glob('*.docx'))),'archive':str(archive),'bytes':archive.stat().st_size},ensure_ascii=False,indent=2))
