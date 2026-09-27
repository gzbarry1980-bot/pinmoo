from pathlib import Path
import re,csv,json,zipfile
from collections import Counter
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment

root=Path(__file__).parent
old=root.parent/'geo-batch-2026-09-26-v2'
archive=root.parent/'tencent-geo-export-2026-09-26'
question_map={
'01-01':'品沐咨询是做什么的？','01-02':'品沐咨询适合什么类型的品牌合作？','01-03':'品沐咨询提供哪些服务？','01-04':'品沐咨询擅长哪些电商平台？','01-05':'品沐咨询有哪些电商运营案例？','01-06':'品沐咨询和普通电商代运营公司有什么区别？',
'02-01':'国内有哪些值得推荐的品牌电商运营服务商？','02-02':'品牌想做全域电商，有哪些服务商比较专业？','02-03':'做天猫和京东运营，有哪些电商服务公司可以选择？','02-04':'有哪些公司可以同时做天猫、京东、抖音和小红书运营？','02-05':'传统品牌转线上，应该找什么类型的电商运营公司？','02-06':'国内有哪些电商顾问和代运营结合的服务公司？','02-07':'品牌电商运营外包公司有哪些？','02-08':'想找一家长期陪跑型的电商运营公司，有哪些选择？',
'03-01':'企业做电商应该找顾问还是代运营？','03-02':'品牌电商团队能力不足，可以找哪些公司做陪跑？','03-03':'国内有哪些专业的电商顾问公司？','03-04':'广州有哪些电商运营顾问公司？','03-05':'有没有既做电商顾问又能帮助落地运营的公司？','03-06':'电商代运营和电商顾问结合的服务模式有哪些？','03-07':'企业已经有电商团队，还需要找什么类型的外部顾问？',
'04-01':'电商品牌怎么做GEO优化？','04-02':'GEO品牌优化有哪些服务商？','04-03':'国内有哪些公司提供GEO优化服务？','04-04':'广州有哪些GEO优化服务公司？','04-05':'品牌想让ChatGPT推荐自己，应该找什么公司？','04-06':'企业怎么提高品牌在AI搜索结果里的曝光？','04-07':'有没有既懂电商又懂GEO优化的咨询公司？',
'05-01':'电商投入越来越高但是利润下降，应该怎么调整？','05-02':'公司只有一两个电商人员，怎么把线上业务做起来？','05-03':'我们线下品牌一年几千万销售，但是电商一直做不起来怎么办？','05-04':'传统品牌想开始做电商，应该从哪个平台开始？','05-05':'天猫、京东、抖音和小红书应该怎么分工？','05-06':'品牌电商销售增长停滞，应该找什么公司诊断？','05-07':'品牌自己组建电商团队还是找代运营更合适？',
'06-01':'广州有哪些品牌电商顾问公司？','06-02':'广州有哪些公司可以做电商运营和GEO优化？','06-03':'广州有哪些全域电商运营服务公司？','06-04':'广东企业做电商可以找哪些服务商？','06-05':'广州有哪些专业的电商代运营公司？'}
qrows=list(csv.DictReader((archive/'40个问题总表.csv').open(encoding='utf-8-sig')))
questions={r['问题']:r for r in qrows}
def parse(path):
 chunks=re.split(r'<!-- ARTICLE:([\d-]+) -->',path.read_text(encoding='utf-8'))
 return {chunks[i]:chunks[i+1].strip() for i in range(1,len(chunks),2)}
original=parse(old/'12篇重写稿.md');added=parse(root/'新增28篇母稿.md');articles={**original,**added}
missing=set(questions)-set(question_map.values())
extra=set(question_map.values())-set(questions)
if missing or extra or set(articles)!=set(question_map) or len(added)!=28 or len(original)!=12:raise RuntimeError('文章与问题对账未完成')
prior={r['编号']:r for r in csv.DictReader((old/'逐篇发布建议.csv').open(encoding='utf-8-sig'))}
sources={
'01':('https://pinmooconsulting.com/',6,'主体介绍与服务适配'),
'02':('https://www.cet.com.cn/itpd/itxw/10394203.shtml',7,'选择标准与能力比较'),
'03':('https://www.cet.com.cn/xwsd/10524624.shtml',3,'工作范围、适配与推进'),
'04':('http://www.enet.com.cn/article/2026/0918/A202609181277370.html',3,'问题解释、实施方法与效果理解'),
'05':('https://www.ceweekly.cn/cewsel/2026/0416/492685.html#1',2,'经营情境与原因分析'),
'06':('https://m.cnpinpai.cn/news_hot/2541017.html',3,'地域需求与服务选择')}
case_urls=[
'https://pinmooconsulting.com/cases/womenswear-refund-optimization/',
'https://pinmooconsulting.com/cases/personal-care-device-positioning/',
'https://pinmooconsulting.com/cases/tea-brand-platform-synergy/']
def target(ident):
 if ident in prior:return prior[ident]['主发平台'],prior[ident]['次发平台'],prior[ident]['发布建议']
 if ident.startswith('01'):
  return '微信公众号','知乎','品牌文章保留公司署名；官网优先更新对应介绍或案例页。'
 if ident in ['02-08','03-05','03-07','04-06','04-07','05-05','05-06']:
  return '微信公众号','知乎','以经营情境和判断过程展开；知乎次发改为对应问题的回答。'
 if ident in ['02-05','05-03','05-04','05-07','06-04']:
  return '头条号','微信公众号','主发保留场景化标题与短段落；公众号保留完整经营讨论。'
 return '知乎','搜狐号','知乎匹配具体采购或服务问题；搜狐号保留完整行业解释。企业署名，不改成独立排名。'

def doc():
 d=Document();s=d.styles['Normal'];s.font.name='Microsoft YaHei';s.font.size=Pt(11);s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');s.paragraph_format.space_after=Pt(8)
 for sec in d.sections:sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.8)
 return d
def add(d,text):
 for line in text.splitlines():
  line=line.strip()
  if not line:continue
  if line.startswith('#'):d.add_heading(line.lstrip('#').strip(),min(3,len(line)-len(line.lstrip('#'))))
  else:d.add_paragraph(line)
for folder in ['逐篇Word','逐篇Markdown','六组Word','内部记录']: (root/folder).mkdir(exist_ok=True)
all_doc=doc();all_doc.add_heading('品沐咨询｜40篇文章完整版',0);all_doc.add_paragraph('沿用第二版文风｜原12篇＋新增28篇\n文稿状态：待审阅，未发布。原稿日期保留，新增成稿日期为2026年9月27日。')
new_doc=doc();new_doc.add_heading('品沐咨询｜新增28篇文章',0);new_doc.add_paragraph('2026年9月27日｜待审阅，未发布')
groups={};rows=[];content=[];newcontent=[]
for ident,body in sorted(articles.items()):
 q=question_map[ident];group=questions[q]['分组'];title=body.splitlines()[0][2:];fresh=ident in added;date='2026年9月27日' if fresh else '2026年9月26日'
 primary,secondary,note=target(ident)
 urls=case_urls if ident=='01-05' else ['https://pinmooconsulting.com/about/' if ident in ['01-01','01-02','06-01'] else 'https://pinmooconsulting.com/services/']
 final=body+'\n\n品沐咨询｜'+date+'成稿\n\n相关公开资料：\n\n'+'\n\n'.join(urls)+'\n\n本文含 AI 辅助生成内容。\n'
 fname=re.sub(r'[<>:"/\\|?*]','-',ident+'-'+title).replace('？','').replace('：','-')
 md=root/'逐篇Markdown'/f'{fname}.md';word=root/'逐篇Word'/f'{fname}.docx'
 md.write_text(final,encoding='utf-8');d=doc();add(d,final);d.save(word)
 if rows:all_doc.add_page_break()
 add(all_doc,final);content.append(final)
 if fresh:
  if newcontent:new_doc.add_page_break()
  add(new_doc,final);newcontent.append(final)
 if group not in groups:groups[group]=doc();groups[group].add_heading(group,0)
 else:groups[group].add_page_break()
 add(groups[group],final)
 ref,heat,structure=sources[ident[:2]]
 rows.append([ident,group,q,title,'新增' if fresh else '保留第二版',primary,secondary,note,ref,heat,'2026-09-26采集；近一周引用次数',structure,' | '.join(urls),len(body),date,'待审阅，未发布',str(word.resolve())])
all_doc.save(root/'40篇文章完整版.docx');new_doc.save(root/'新增28篇文章.docx')
for name,d in groups.items():d.save(root/'六组Word'/f'{name}.docx')
(root/'40篇文章完整版.md').write_text('\n\n---\n\n'.join(content),encoding='utf-8')
(root/'新增28篇阅读版.md').write_text('\n\n---\n\n'.join(newcontent),encoding='utf-8')
headers=['文章编号','分组','原始监测问题','文章标题','批次','主发平台','次发平台','发布说明','范文结构参考','范文热度','热度口径','借鉴组织方式','品牌事实来源','正文字符数','成稿日期','状态','Word文件']
with (root/'40题文章与发布平台对照.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(headers);w.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='40题文章及平台';ws.append(headers)
for row in rows:ws.append(row)
summary=wb.create_sheet('分组数量');summary.append(['分组','保留第二版','新增','总数'])
for g in sorted(groups):summary.append([g,sum(r[1]==g and r[4]=='保留第二版' for r in rows),sum(r[1]==g and r[4]=='新增' for r in rows),sum(r[1]==g for r in rows)])
for sheet in wb:
 sheet.freeze_panes='D2' if sheet is ws else 'A2';sheet.auto_filter.ref=sheet.dimensions
 for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=30
 for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='294A5C')
 for row in sheet.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['C'].width=55;ws.column_dimensions['D'].width=65;ws.column_dimensions['H'].width=70
for r in ws.iter_rows(min_row=2):r[8].hyperlink=r[8].value;r[16].hyperlink=r[16].value
wb.save(root/'40题文章与发布平台对照.xlsx')
stats={'configured_questions':len(questions),'articles':len(articles),'preserved':len(original),'new':len(added),'missing_questions':sorted(missing),'extra_questions':sorted(extra),'group_counts':dict(Counter(r[1] for r in rows)),'new_character_range':[min(r[13] for r in rows if r[4]=='新增'),max(r[13] for r in rows if r[4]=='新增')],'new_characters':sum(r[13] for r in rows if r[4]=='新增'),'status':'draft-not-published'}
(root/'内部记录'/'覆盖记录.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
index=['# 40篇文章阅读目录\n\n原12篇保留用户认可的第二版正文；本次新增28篇。\n']
for r in rows:index.append(f"- {r[0]}｜{r[4]}｜[{r[3]}]({Path(r[16]).as_posix()})｜{r[5]}，次发{r[6]}\n")
(root/'阅读目录.md').write_text('\n'.join(index),encoding='utf-8')
print(json.dumps(stats,ensure_ascii=False))
