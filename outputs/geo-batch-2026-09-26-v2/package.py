from pathlib import Path
import re,csv,zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment

p=Path(__file__).parent
source=p/'12篇重写稿.md'
parts=re.split(r'<!-- ARTICLE:([\d-]+) -->',source.read_text(encoding='utf-8'))
groups={'01':'品沐品牌认知','02':'电商服务商推荐','03':'电商顾问与代运营','04':'GEO与AI搜索优化','05':'品牌经营需求','06':'广州广东区域'}
destinations={
'01-01':('微信公众号','知乎','用公司介绍长文发布；官网已有关于页吸收相关段落，不再复制一篇近似介绍。'),
'01-02':('微信公众号','知乎','面向已有团队和传统品牌；知乎选择适合对象、顾问价值类问题。'),
'02-01':('知乎','搜狐号','适配电商服务商选择类问题；保留品沐署名及文中自身立场。'),
'02-02':('微信公众号','搜狐号','用多平台协作困境吸引品牌负责人；次发可改标题为“全域电商做了多个平台，生意为什么仍未衔接？”'),
'03-01':('知乎','微信公众号','适配自营、顾问与代运营选择问题；开头直接保留店铺困境。'),
'03-02':('微信公众号','头条号','面向品牌管理者与运营负责人；次发突出周会反复讨论却难落地的场景。'),
'04-01':('微信公众号','知乎','面向已听说GEO但不理解内容差异的老板；知乎适配品牌为何不被AI推荐的问题。'),
'04-02':('知乎','搜狐号','作为GEO采购讨论文章；不包装为第三方排名或评测。'),
'05-01':('微信公众号','头条号','突出活动后利润落差，保留后文的退款周期与成本解释。'),
'05-02':('头条号','微信公众号','用小团队工作压力切入，适合传统企业负责人阅读。'),
'06-01':('知乎','搜狐号','广州电商顾问选择类问题；不加“十大”“第一”等未经支持的标题。'),
'06-02':('微信公众号','搜狐号','解释电商业务与GEO怎样协作；次发可用“广州品牌如何协同电商经营与GEO内容？”')}
def initdoc():
 d=Document();s=d.styles['Normal'];s.font.name='Microsoft YaHei';s.font.size=Pt(11);s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');s.paragraph_format.space_after=Pt(8);return d
def insert(d,text):
 for line in text.splitlines():
  line=line.strip()
  if not line:continue
  if line.startswith('#'):d.add_heading(line.lstrip('#').strip(),level=min(3,len(line)-len(line.lstrip('#'))))
  else:d.add_paragraph(line)
(p/'逐篇Word').mkdir(exist_ok=True);(p/'逐篇Markdown').mkdir(exist_ok=True)
whole=initdoc();whole.add_heading('品沐咨询｜12篇重写稿',0);whole.add_paragraph('第二版 · 2026年9月26日\n正文供审阅，尚未发布。逐篇发布建议另附。')
rows=[];alltext=[]
for i in range(1,len(parts),2):
 ident=parts[i];body=parts[i+1].strip();title=body.splitlines()[0][2:];group=groups[ident[:2]];primary,secondary,note=destinations[ident]
 url='https://pinmooconsulting.com/about/' if ident in ['01-01','01-02','06-01'] else 'https://pinmooconsulting.com/services/'
 final=body+'\n\n品沐咨询｜2026年9月26日成稿\n\n品牌相关介绍可参阅品沐官网：'+url+'\n\n本文含 AI 辅助生成内容。\n'
 filename=ident+'-'+title.replace('？','').replace('：','-')
 (p/'逐篇Markdown'/f'{filename}.md').write_text(final,encoding='utf-8')
 d=initdoc();insert(d,final);d.save(p/'逐篇Word'/f'{filename}.docx')
 if rows:whole.add_page_break()
 insert(whole,final);alltext.append(final)
 rows.append([ident,group,title,primary,secondary,note,len(body),'待审阅，未发布'])
whole.save(p/'12篇重写稿-阅读版.docx')
(p/'12篇重写稿-阅读版.md').write_text('\n\n---\n\n'.join(alltext),encoding='utf-8')
headers=['编号','分组','文章标题','主发平台','次发平台','发布建议','正文字符数','状态']
with (p/'逐篇发布建议.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(headers);w.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='逐篇发布建议';ws.append(headers)
for row in rows:ws.append(row)
ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=35
ws.column_dimensions['C'].width=62;ws.column_dimensions['F'].width=70
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='294A5C')
for row in ws.iter_rows(min_row=2):
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
wb.save(p/'逐篇发布建议.xlsx')
old=p.parent/'geo-batch-2026-09-26'/'README.md'
previous=old.read_text(encoding='utf-8')
marker='> **旧版弃用：本批稿件已因文风与可读性问题被否定，请勿按此版本发布。重写稿见相邻目录 geo-batch-2026-09-26-v2。**\n\n'
if not previous.startswith(marker):old.write_text(marker+previous,encoding='utf-8')
(p.parent/'geo-batch-2026-09-26'/'旧版弃用说明.txt').write_text('本目录第一版文章已弃用，不作为发布稿。请使用 geo-batch-2026-09-26-v2 内的重写稿；该版仍待用户审阅，未发布。',encoding='utf-8')
print({'articles':len(rows),'characters':[r[6] for r in rows]})
